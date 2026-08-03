"""Hybrid Excel parser: deterministic heavy-lifting + optional LLM inference.

The parser tries everything it can without an LLM (so the app still works
before a GEMINI_API_KEY is set). For genuinely messy sheets we ask the LLM
to infer/clean the structure.
"""
from __future__ import annotations

import io
import re
from collections import Counter
from typing import Any, List, Optional

import pandas as pd
from openpyxl import load_workbook

from .llm import llm_available, llm_json


# --------------------------------------------------------------------------- #
# Sheet discovery
# --------------------------------------------------------------------------- #
def list_sheets(file_bytes: bytes, filename: str) -> List[dict]:
    sheets = []
    if filename.lower().endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        for name in wb.sheetnames:
            ws = wb[name]
            sheets.append({"name": name, "dims": ws.dimensions})
    else:
        xl = pd.ExcelFile(io.BytesIO(file_bytes))
        for name in xl.sheet_names:
            sheets.append({"name": name, "dims": ""})
    return sheets


def read_sheet(file_path: str, sheet_name: str) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name)


def read_and_locate_header(file_path: str, sheet_name: str) -> pd.DataFrame:
    """Read a sheet handling title/metadata rows above the real header.

    Reads with header=None, drops empty rows, then picks the earliest row among
    the first 12 with the most non-empty cells as the column header. This handles
    the classic "title row above the table" messy-Excel case.
    """
    raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
    raw = raw.replace("", None).dropna(how="all").reset_index(drop=True)
    if raw.empty:
        return pd.DataFrame()

    n_rows = min(12, len(raw))
    best_idx, best_count = 0, -1
    for i in range(n_rows):
        count = int(raw.iloc[i].notna().sum())
        if count > best_count:
            best_count, best_idx = count, i

    header = [str(v).strip() if pd.notna(v) else f"col_{j}" for j, v in enumerate(raw.iloc[best_idx].tolist())]
    body = raw.iloc[best_idx + 1 :].copy()
    body.columns = _dedupe_names(header)
    return body.reset_index(drop=True)


# --------------------------------------------------------------------------- #
# Cleaning heuristics (no LLM required)
# --------------------------------------------------------------------------- #
def _dedupe_names(names: List[str]) -> List[str]:
    counts: Counter = Counter(names)
    used: Counter = Counter()
    out = []
    for n in names:
        base = str(n).strip()
        if counts[base] > 1:
            used[base] += 1
            out.append(f"{base}_{used[base]}")
        else:
            out.append(base)
    return out


def clean_dataframe(df: pd.DataFrame) -> tuple[pd.DataFrame, List[str]]:
    notes: List[str] = []
    df = df.copy()

    before = df.shape
    df = df.dropna(axis=0, how="all").dropna(axis=1, how="all")
    if df.shape != before:
        notes.append(
            f"Dropped {before[0]-df.shape[0]} fully-empty rows and {before[1]-df.shape[1]} fully-empty columns."
        )

    df.columns = [str(c).strip() for c in df.columns]
    df.columns = _dedupe_names(list(df.columns))

    for c in df.select_dtypes(include=["object"]).columns:
        df[c] = df[c].astype(str).str.strip()
        df[c] = df[c].replace({"": None, "nan": None})

    for c in df.columns:
        if _looks_numeric(df[c]):
            cleaned = df[c].astype(str).str.replace(",", "", regex=False).str.replace("$", "", regex=False)
            df[c] = pd.to_numeric(cleaned, errors="coerce")
            notes.append(f"Column '{c}' coerced to numeric.")
        elif _looks_date(df[c]):
            try:
                df[c] = pd.to_datetime(df[c], errors="coerce")
                notes.append(f"Column '{c}' parsed as datetime.")
            except Exception:
                pass

    df = df.dropna(axis=1, how="all")
    return df, notes


def _looks_numeric(s: pd.Series) -> bool:
    if s.dtype.kind in "ifu":
        return False
    sample = s.dropna().astype(str).str.strip().head(50)
    if len(sample) == 0:
        return False
    cleaned = sample.str.replace(",", "", regex=False).str.replace("$", "", regex=False)
    try:
        pd.to_numeric(cleaned)
        return True
    except Exception:
        return False


def _looks_date(s: pd.Series) -> bool:
    sample = s.dropna().astype(str).head(20)
    if len(sample) == 0:
        return False
    if sample.str.match(r"^\d{4}-\d{2}-\d{2}$").all() or sample.str.match(r"^\d{1,2}/\d{1,2}/\d{2,4}$").all():
        return True
    return False


# --------------------------------------------------------------------------- #
# LLM-assisted structure inference (messy / ambiguous sheets)
# --------------------------------------------------------------------------- #
def llm_infer_structure(file_path: str, sheet_name: str) -> Optional[pd.DataFrame]:
    if not llm_available():
        return None
    try:
        raw = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
    except Exception:
        return None
    preview = _render_compact(raw.head(40))

    prompt = f"""You are a meticulous data-cleaner. Below is a JSON array of raw rows
extracted from an Excel sheet named '{sheet_name}'. It may contain merged cells,
notes, multiple table blocks, or no meaningful headers.

Your job:
1. Identify the single most important TABLE of data in these rows.
2. Determine the best header names (the first meaningful row becomes column names).
3. Output the data as a JSON object: {{"columns":[...], "rows":[[...],[...]]}}.
   - Drop title/note/blank rows.
   - Coerce numbers to numbers, dates to ISO 'YYYY-MM-DD'.
   - Replace empty cells with null.
Keep at most 25 rows. If unsure, still return your best guess.

RAW DATA:
{preview}
"""
    result = llm_json(prompt)
    if not result:
        return None
    try:
        cols = result["columns"]
        rows = result["rows"]
        return pd.DataFrame(rows, columns=cols)
    except Exception:
        return None


def _render_compact(df: pd.DataFrame) -> str:
    lines = [str(df.columns.tolist())]
    for _, row in df.head(40).iterrows():
        lines.append(row.tolist().__str__())
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
# Metadata helpers
# --------------------------------------------------------------------------- #
def preview_frame(df: pd.DataFrame, limit: int = 8) -> dict:
    cols = [str(c) for c in df.columns]
    rows = df.head(limit).where(pd.notna(df), None).values.tolist()
    return {"columns": cols, "rows": rows}


def summarize_columns(df: pd.DataFrame) -> List[dict]:
    out = []
    for c in df.columns:
        s = df[c]
        if pd.api.types.is_numeric_dtype(s):
            role = "measure"
        elif pd.api.types.is_datetime64_any_dtype(s):
            role = "date"
        else:
            role = "dimension"
        out.append(
            {
                "name": str(c),
                "dtype": str(s.dtype),
                "nullable": int(s.isna().sum()) > 0,
                "description": _readable_name(str(c)),
                "role": role,
            }
        )
    return out


def _readable_name(col: str) -> str:
    tokens = [t for t in re.split(r"[^A-Za-z0-9]+", col) if t]
    return " ".join(t.capitalize() for t in tokens) or col


def infer_quality(df: pd.DataFrame) -> tuple[str, List[str]]:
    notes: List[str] = []
    total = len(df)
    if total == 0:
        return "empty", ["No rows detected after cleaning."]
    missing = int(df.isna().sum().sum())
    missing_pct = missing / (total * df.shape[1])
    if missing_pct == 0:
        quality = "excellent"
        notes.append("No missing values detected.")
    elif missing_pct < 0.05:
        quality = "good"
        notes.append(f"{missing_pct:.1%} of cells are empty.")
    elif missing_pct < 0.2:
        quality = "fair"
        notes.append(f"{missing_pct:.1%} of cells are empty.")
    else:
        quality = "poor"
        notes.append(f"{missing_pct:.1%} of cells are empty; consider cleaning.")
    return quality, notes