"""Smoke test: upload a messy Excel, check pipeline + dashboard fallbacks (no LLM key)."""
import io
import sys
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parent))
from app.main import app


def make_messy_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Quarterly Sales Report 2024"  # title row (should be dropped)
    ws["A2"] = "Region"
    ws["B2"] = "Sales ($)"
    ws["C2"] = "Month"
    ws.append(["North", 1200, "2024-01"])
    ws.append(["South", 3400, "2024-01"])
    ws.append(["North", 2600, "2024-02"])
    ws.append(["South", 2900, "2024-02"])
    ws.append(["", "", ""])  # empty row (dropped)
    ws.append(["East", "1,500", "2024-03"])  # messy number
    ws.append(["West", 1800, "March 2024"])
    ws.append(["East", 2000, "2024-04"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def test_flow():
    data = make_messy_xlsx()
    client = TestClient(app)
    r = client.post("/api/datasets/upload", files={"file": ("sales.xlsx", data, "application/vnd...xlsx")})
    assert r.status_code == 200, r.text
    d = r.json()
    ds = d["dataset_id"]
    assert d["row_count"] > 0
    print("UPLOAD OK:", d["quality"], "| rows:", d["row_count"], "| notes:", d["notes"])

    sin = client.post(f"/api/datasets/{ds}/insights")
    assert sin.status_code == 200, sin.text
    doc = sin.json()
    print("INSIGHTS OK, summary[:80]:", doc["summary"][:80])

    su = client.post(f"/api/datasets/{ds}/dashboards/suggest")
    assert su.status_code == 200, su.text
    suggs = su.json()["suggestions"]
    print("SUGGEST OK:", [s["title"] for s in suggs])

    b = client.post(f"/api/datasets/{ds}/dashboards/build", json={"suggestion_id": suggs[0]["id"]})
    assert b.status_code == 200, b.text
    spec = b.json()
    print("BUILD OK, layouts:", len(spec["layouts"]))

    c = client.post(f"/api/datasets/{ds}/chat", json={"question": "total sales by region"})
    assert c.status_code == 200, c.text
    print("CHAT OK:", c.json()["answer"][:60])

    dl = client.get(f"/api/datasets/{ds}/download")
    assert dl.status_code == 200
    print("DOWNLOAD OK:", len(dl.content), "bytes")
    print("\nALL PASSED")


if __name__ == "__main__":
    test_flow()